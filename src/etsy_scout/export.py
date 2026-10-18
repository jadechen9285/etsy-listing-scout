"""Write ListingRows to Excel (sheet `competitor_listing`) and CSV.

A multi-keyword batch is written as one combined file; its Excel copy also gets a
`keyword_summary` sheet with one row per keyword.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from .models import COLUMNS, ListingRow

SHEET_NAME = "competitor_listing"
SUMMARY_SHEET = "keyword_summary"
EXCEL_MAX_ROWS = 1_048_575  # Excel's sheet limit, minus the header row


class TooBigForExcel(ValueError):
    pass


def to_dataframe(rows: list[ListingRow]) -> pd.DataFrame:
    return pd.DataFrame([r.to_record() for r in rows], columns=COLUMNS)


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_") or "search"


def summary_frame(df: pd.DataFrame) -> pd.DataFrame:
    """One row per keyword: how many listings, typical price, sales, and popularity."""
    if df.empty:
        return pd.DataFrame()
    grouped = df.groupby("target_keyword", sort=False)
    out = pd.DataFrame(
        {
            "listings": grouped.size(),
            "rank_method": grouped["rank_method"].first(),
            "median_price_dollar": grouped["price_dollar"].median().round(2),
            "min_price_dollar": grouped["price_dollar"].min(),
            "max_price_dollar": grouped["price_dollar"].max(),
            "share_on_sale_pct": (
                grouped["on_sales_price_dollar"].apply(lambda s: s.notna().mean() * 100).round(1)
            ),
            "median_num_favorers": grouped["num_favorers"].median(),
            "median_reviews": grouped["reviews"].median(),
            "median_shop_total_sales": grouped["shop_total_sales"].median(),
        }
    )
    return out.reset_index()


def excel_bytes(df: pd.DataFrame, summary: pd.DataFrame | None = None) -> bytes:
    if len(df) > EXCEL_MAX_ROWS:
        raise TooBigForExcel(f"{len(df):,} rows is more than Excel allows ({EXCEL_MAX_ROWS:,}).")
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        if summary is not None and not summary.empty:
            summary.to_excel(writer, sheet_name=SUMMARY_SHEET, index=False)
            sws = writer.sheets[SUMMARY_SHEET]
            sws.freeze_panes = "B2"
            for i, col in enumerate(summary.columns, start=1):
                sws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 2)
        df.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        ws = writer.sheets[SHEET_NAME]
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, start=1):
            width = 60 if col == "listing_link" else max(12, min(40, len(col) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
        link_col = df.columns.get_loc("listing_link") + 1
        for r in range(2, len(df) + 2):
            cell = ws.cell(row=r, column=link_col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    return buf.getvalue()


def batch_name(keywords: list[str]) -> str:
    """File-name stem for a run: the keyword, or the first keyword plus a count."""
    if len(keywords) == 1:
        return keywords[0]
    return f"{keywords[0]} and {len(keywords) - 1} more"


def write_outputs(
    rows: list[ListingRow], name: str, out_dir: Path, summary: bool = False
) -> tuple[Path | None, Path]:
    """Write `<name>_<timestamp>.xlsx` and `.csv`. Returns (xlsx or None, csv).

    With `summary=True` (multi-keyword batches) the Excel file gets a keyword_summary
    sheet. If the rows don't fit in an Excel sheet, only the CSV is written.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{slugify(name)}_{datetime.now():%Y%m%d_%H%M%S}"
    df = to_dataframe(rows)
    csv = out_dir / f"{stem}.csv"
    df.to_csv(csv, index=False)
    xlsx: Path | None = out_dir / f"{stem}.xlsx"
    try:
        xlsx.write_bytes(excel_bytes(df, summary_frame(df) if summary else None))
    except TooBigForExcel:
        xlsx = None
    return xlsx, csv
